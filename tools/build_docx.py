"""ממיר את docs/grammar.md לקובץ Word (‏.docx) מימין לשמאל.

שימוש: python tools/build_docx.py [docs/grammar.md] [milon-klei-kodesh-grammar.docx]
דורש: pip install python-docx
"""
import json, os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

SRC, OUT = (sys.argv[1:3] + ["docs/grammar.md", "milon-klei-kodesh-grammar.docx"][len(sys.argv[1:3]):])[:2]
FONT = "David"
BODY_PT = 13
GH = "https://github.com/Hadabadra/KLI-MULON/releases"
STORE = "https://otzaria.org/plugins"
LABELS = {"קצר": "בקצרה:", "מקור": "מקור:", "דוגמה": "דוגמה מן המקרא:"}
FIGS = {"moza": "מוצאות הפה", "nikud": "סימני הניקוד", "syl": "חלוקה להברות",
        "ranks": "דרגות המפסיקים", "tree": "חלוקת הפסוק לפי הטעמים"}
LATIN = re.compile(r"[A-Za-z][A-Za-z0-9:/._\-?=&%#+]*(?:[ ,/]+[A-Za-z0-9:/._\-?=&%#+]*[A-Za-z0-9/])*")


# ---------- עזרי XML ----------
def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def _set(parent, tag, **attrs):
    """מחליף/יוצר רכיב יחיד בשם tag תחת parent."""
    for old in parent.findall(qn(tag)):
        parent.remove(old)
    e = _el(tag, **attrs)
    parent.append(e)
    return e


def set_fonts(rpr, size_pt=None, bold=None):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = _el("w:rFonts")
        rpr.insert(0, rf)
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(qn("w:" + a), None)
    for a in ("ascii", "hAnsi", "cs"):
        rf.set(qn("w:" + a), FONT)
    if size_pt:
        for t in ("w:sz", "w:szCs"):
            _set(rpr, t, val=int(size_pt * 2))
    if bold:
        for t in ("w:b", "w:bCs"):
            _set(rpr, t)
    return rpr


def bidi_par(p, align=None):
    ppr = p._p.get_or_add_pPr()
    if ppr.find(qn("w:bidi")) is None:
        ppr.append(_el("w:bidi"))
    if align is not None:
        p.alignment = align
    return p


def add_run(p, text, bold=False, italic=False, size=None, rtl=True):
    r = p.add_run(text)
    rpr = r._r.get_or_add_rPr()
    set_fonts(rpr)
    if bold:
        r.bold = True
        rpr.append(_el("w:bCs"))
    if italic:
        r.italic = True
        rpr.append(_el("w:iCs"))
    if size:
        r.font.size = Pt(size)
        rpr.append(_el("w:szCs", val=int(size * 2)))
    if rtl:
        rpr.append(_el("w:rtl"))
    return r


def add_text(p, text, bold=False, italic=False, size=None):
    """מוסיף טקסט: קטעי עברית כ-run מימין לשמאל, וקטעי אותיות לטיניות (a, ts, כתובות) בלי rtl."""
    pos = 0
    for m in LATIN.finditer(text):
        if m.start() > pos:
            add_run(p, text[pos:m.start()], bold, italic, size)
        add_run(p, m.group(), bold, italic, size, rtl=False)
        pos = m.end()
    if pos < len(text):
        add_run(p, text[pos:], bold, italic, size)


def add_hyperlink(p, url, text):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = _el("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = _el("w:r")
    rpr = _el("w:rPr")
    set_fonts(rpr)
    rpr.append(_el("w:color", val="0563C1"))
    rpr.append(_el("w:u", val="single"))
    r.append(rpr)
    t = _el("w:t")
    t.text = text
    r.append(t)
    h.append(r)
    p._p.append(h)


def _fc(kind):
    r = _el("w:r")
    r.append(_el("w:fldChar", fldCharType=kind))
    return r


def add_field(p, instr, placeholder, end=True):
    fc = _fc
    p._p.append(fc("begin"))
    r = _el("w:r")
    it = _el("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r.append(it)
    p._p.append(r)
    p._p.append(fc("separate"))
    add_run(p, placeholder)
    if end:
        p._p.append(fc("end"))


# ---------- סדר הרכיבים לפי הסכמה (Word מסרב לפתוח קובץ שהסדר בו שגוי) ----------
ORDER = {
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd "
           "snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment "
           "textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs "
           "em lang eastAsianLayout specVanish oMath",
    "sectPr": "headerReference footerReference footnotePr endnotePr type pgSz pgMar paperSrc pgBorders lnNumType "
              "pgNumType cols formProt vAlign noEndnote titlePg textDirection bidi rtlGutter docGrid printerSettings",
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
}
ORDER = {qn("w:" + k): {qn("w:" + t): n for n, t in enumerate(v.split())} for k, v in ORDER.items()}


def normalize(root):
    for el in root.iter(*ORDER):
        rank = ORDER[el.tag]
        kids = list(el)
        kids.sort(key=lambda c: rank.get(c.tag, len(rank)))
        el[:] = kids


# ---------- תוכן מוטבע ----------
LINK = re.compile(r"\[\[([^\]|]+)(?:\|([^\]]+))?\]\]")


def link_text(m):
    book, term = m.group(1).strip(), (m.group(2) or "").strip()
    if not term:
        return book
    if book == "חיפוש":
        return term
    if book == "פסוק":
        return f"({term})"
    return f"{book}: {term}"


def inline(p, text, bold=False):
    text = LINK.sub(link_text, text)
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            add_text(p, part, bold=bold or i % 2 == 1)


# ---------- סגנונות ----------
def setup_styles(doc):
    st = doc.styles
    # ברירת מחדל של המסמך: גופן עברי, שפה עברית, מימין לשמאל
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_fonts(rpr_default, BODY_PT)
    _set(rpr_default, "w:lang", val="en-US", bidi="he-IL")
    for name, size, bold in (("Normal", BODY_PT, False), ("Title", 28, True), ("Subtitle", 16, False),
                             ("Heading 1", 20, True), ("Heading 2", 16, True), ("Heading 3", 14, True),
                             ("List Bullet", BODY_PT, False), ("List Paragraph", BODY_PT, False),
                             ("TOC Heading", 18, True)):
        try:
            s = st[name]
        except KeyError:
            continue
        rpr = s.element.get_or_add_rPr()
        set_fonts(rpr, size, bold)
        if rpr.find(qn("w:rtl")) is None:
            rpr.append(_el("w:rtl"))
        ppr = s.element.get_or_add_pPr()
        if ppr.find(qn("w:bidi")) is None:
            ppr.append(_el("w:bidi"))
    for name, col in (("Heading 1", "1F3864"), ("Heading 2", "2E5496"), ("Heading 3", "2E5496"), ("Title", "1F3864")):
        st[name].font.color.rgb = RGBColor.from_string(col)
    st["Normal"].paragraph_format.space_after = Pt(4)
    st["Normal"].paragraph_format.line_spacing = 1.15
    st["Heading 1"].paragraph_format.page_break_before = True
    st["Heading 1"].paragraph_format.space_after = Pt(10)
    st["Heading 2"].paragraph_format.space_before = Pt(16)
    st["Heading 3"].paragraph_format.space_before = Pt(10)
    # תוכן העניינים: גופן עברי ומימין לשמאל גם לרמות שייווצרו בעדכון השדה
    for lvl in (1, 2, 3):
        try:
            s = st[f"toc {lvl}"]
        except KeyError:  # הסגנון המובנה של Word ‏(toc 1..3), ולא סגנון מותאם בשם דומה
            s = st.add_style(f"toc {lvl}", 1)
            s.element.attrib.pop(qn("w:customStyle"), None)
            s.style_id = f"TOC{lvl}"
            s.base_style = st["Normal"]
            s.element.get_or_add_next().val = "Normal"
        set_fonts(s.element.get_or_add_rPr(), BODY_PT, lvl == 1)
        ppr = s.element.get_or_add_pPr()
        if ppr.find(qn("w:bidi")) is None:
            ppr.append(_el("w:bidi"))
        s.paragraph_format.left_indent = Cm(0.6 * (lvl - 1))  # w:left בפסקה מימין לשמאל = תחילת השורה (מימין)
        s.paragraph_format.space_after = Pt(2)


def setup_section(sec):
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    if sec._sectPr.find(qn("w:bidi")) is None:
        sec._sectPr.append(_el("w:bidi"))
    fp = bidi_par(sec.footer.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
    add_field(fp, "PAGE", "1")


# ---------- טבלה ----------
def add_table(doc, rows):
    rows = [r for r in rows if not re.fullmatch(r"[\s|:\-]*", "|".join(r))]
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblpr = t._tbl.tblPr
    tblpr.append(_el("w:bidiVisual"))  # עמודה ראשונה מימין
    _set(tblpr, "w:tblW", w=5000, type="pct")
    for i, row in enumerate(rows):
        tr = t.rows[i]
        if i == 0:
            trpr = tr._tr.get_or_add_trPr()
            trpr.append(_el("w:tblHeader"))  # חוזרת בראש כל עמוד
        for j in range(ncol):
            cell = tr.cells[j]
            p = bidi_par(cell.paragraphs[0])
            p.paragraph_format.space_after = Pt(0)
            inline(p, row[j] if j < len(row) else "", bold=(i == 0))
            if i == 0:
                cell._tc.get_or_add_tcPr().append(_el("w:shd", val="clear", color="auto", fill="D9E2F3"))
    doc.add_paragraph()


# ---------- המרה ----------
def load_meta():
    try:
        s = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "grammar.js"), encoding="utf8").read()
        o = json.loads(s[s.index("{"):s.rstrip().rstrip(";").rindex("}") + 1])
        return o.get("title", ""), o.get("note", "")
    except Exception:
        return "", ""


def title_page(doc):
    sub, note = load_meta()
    for _ in range(5):
        bidi_par(doc.add_paragraph())
    p = bidi_par(doc.add_paragraph(style="Title"), WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, "ספר הדקדוק — לשון הקודש")
    if sub:
        p = bidi_par(doc.add_paragraph(style="Subtitle"), WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, sub)
    bidi_par(doc.add_paragraph())
    p = bidi_par(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, 'מתוך התוסף "לשון הקודש" (מילון כליקודש) לתוכנת אוצריא.', size=12)
    if note:
        p = bidi_par(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, note, italic=True, size=11)
    bidi_par(doc.add_paragraph())
    p = bidi_par(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, "הערה: ", bold=True, size=12)
    add_text(p, "בתוסף עצמו יש ציורים, קישורים פעילים לספרי הספרייה וחיפוש. להורדת התוסף:", size=12)
    for label, url in (("GitHub (דף ההורדות): ", GH), ("חנות התוספים של אוצריא: ", STORE)):
        p = bidi_par(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, label, size=12)
        add_hyperlink(p, url, url)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def toc(doc, lines):
    p = bidi_par(doc.add_paragraph(style="TOC Heading" if "TOC Heading" in [s.name for s in doc.styles] else "Heading 1"))
    p.paragraph_format.page_break_before = False
    add_text(p, "תוכן העניינים")
    # השדה מוצג מראש עם רשימת החלקים והפרקים (בלי מספרי עמודים), כדי שיהיה תוכן עניינים
    # גם בתוכנה שאינה מעדכנת שדות; Word מעדכן אותו בפתיחה ומוסיף עמודים וקישורים.
    p = bidi_par(doc.add_paragraph())
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "(לעדכון מספרי העמודים: לחיצה ימנית ← עדכן שדה, או F9)", end=False)
    for l in lines:
        m = re.match(r"(#{1,2})\s+(.*)", l)
        if m:
            p = bidi_par(doc.add_paragraph(style=f"toc {len(m.group(1))}"))
            add_text(p, re.sub(r"\*\*", "", LINK.sub(link_text, m.group(2))))
    p._p.append(_fc("end"))
    # מבקש מ-Word לעדכן את השדות בפתיחה
    settings = doc.settings.element
    for old in settings.findall(qn("w:updateFields")):
        settings.remove(old)
    uf = _el("w:updateFields", val="true")
    after = [qn("w:" + t) for t in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                                       "mathPr", "attachedSchema", "themeFontLang", "clrSchemeMapping",
                                       "doNotIncludeSubdocsInStats", "doNotAutoCompressPictures", "forceUpgrade",
                                       "captions", "readModeInkLockDown", "smartTagType", "schemaLibrary",
                                       "shapeDefaults", "doNotEmbedSmartTags", "decimalSymbol", "listSeparator")]
    nxt = next((c for c in settings if c.tag in after), None)
    if nxt is not None:
        nxt.addprevious(uf)
    else:
        settings.append(uf)


def convert(doc, lines):
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip("\n")
        s = raw.strip()
        i += 1
        if not s:
            continue
        m = re.match(r"(#{1,3})\s+(.*)", s)
        if m:
            p = bidi_par(doc.add_heading(level=len(m.group(1))))
            inline(p, m.group(2))
            continue
        if s.startswith("|"):
            rows = [s]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            add_table(doc, [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows])
            continue
        m = re.match(r"!ציור:\s*(\S+)", s)
        if m:
            name = FIGS.get(m.group(1), m.group(1))
            p = bidi_par(doc.add_paragraph(), WD_ALIGN_PARAGRAPH.CENTER)
            add_text(p, f"[ציור: {name} — ראו בתוסף]", italic=True, size=11)
            continue
        m = re.match(r":(קצר|מקור|דוגמה):\s*(.*)", s)
        if m:
            p = bidi_par(doc.add_paragraph())
            add_text(p, LABELS[m.group(1)] + " ", bold=True)
            inline(p, m.group(2))
            continue
        if s.startswith("- "):
            p = bidi_par(doc.add_paragraph(style="List Bullet"))
            inline(p, s[2:])
            continue
        p = bidi_par(doc.add_paragraph())
        m = re.match(r"(\d+\.)\s+(.*)", s)
        if m:
            p.paragraph_format.left_indent = Cm(0.6)  # תחילת השורה
            add_text(p, m.group(1) + " ")
            inline(p, m.group(2))
            continue
        m = re.match(r"(ראו גם:)\s*(.*)", s)
        if m:
            add_text(p, m.group(1) + " ", italic=True)
            inline(p, m.group(2))
            continue
        inline(p, s)


def main():
    lines = open(SRC, encoding="utf8").read().splitlines()
    doc = Document()
    setup_styles(doc)
    setup_section(doc.sections[0])
    doc.core_properties.title = "ספר הדקדוק — לשון הקודש"
    doc.core_properties.language = "he-IL"
    title_page(doc)
    toc(doc, lines)
    # ההקדמה (לפני הכותרת הראשונה) באה מיד אחרי תוכן העניינים
    first = next((k for k, l in enumerate(lines) if l.startswith("# ")), len(lines))
    intro, body = lines[:first], lines[first:]
    if any(l.strip() for l in intro):
        p = bidi_par(doc.add_paragraph())
        p.add_run().add_break(WD_BREAK.PAGE)
        convert(doc, intro)
    convert(doc, body)
    # הפסקה הריקה שיוצר python-docx בתחילת מסמך חדש — גם היא מימין לשמאל
    for p in doc.paragraphs:
        bidi_par(p)
    for part in (doc.element, doc.styles.element, doc.sections[0].footer._element):
        normalize(part)
    doc.save(OUT)
    print("ok", OUT)


if __name__ == "__main__":
    main()
